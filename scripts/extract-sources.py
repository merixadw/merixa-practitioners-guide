"""
Page-limited PDF/DOCX text extraction for Merixa Practitioner's Guide ingest.
Writes plain-text extracts under content/raw-cache for the Node card builder.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from concurrent.futures import ThreadPoolExecutor
from concurrent.futures import TimeoutError as FuturesTimeout
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
DOWNLOADS = ROOT.parent
CACHE = ROOT / "content" / "raw-cache"
REPORT = ROOT / "content" / "extract-report.json"

SOURCE_ROOTS = [
    DOWNLOADS / "Merixa x Codex" / "IFRS",
    DOWNLOADS / "Merixa x Codex" / "Management Reporting",
    DOWNLOADS / "Merixa x Codex" / "Financial Reporting",
    DOWNLOADS / "Merixa x Codex" / "Financial Management",
    DOWNLOADS / "Merixa x Codex" / "Audit",
    DOWNLOADS / "Merixa x Codex" / "Strategic Management",
    DOWNLOADS / "Merixa x Codex" / "Project Management",
    DOWNLOADS / "Merixa_Management_Reporting_Practitioner_Glossary_v2_QA_Rebuilt_Pack",
]

SKIP_DIRS = {
    ".git",
    ".venv",
    "node_modules",
    ".next",
    "dist",
    "__pycache__",
    ".claude",
    ".claude-flow",
    ".swarm",
    "mpl-data",
    "site-packages",
    "merixa_internal_delivery_platform",
    "Merixa_Finance_ETL",
    "Merixa Analysing Tool",
    "exports",
    "P&O Ferries",
}

PRIORITY_PREFIXES = (
    "Merixa x Codex/IFRS",
    "Merixa x Codex/Management Reporting",
    "Merixa x Codex/Financial Reporting",
    "Merixa x Codex/Financial Management",
    "Merixa x Codex/Audit",
    "Merixa x Codex/Strategic Management",
    "Merixa x Codex/Project Management",
    "Merixa_Management_Reporting_Practitioner_Glossary_v2_QA_Rebuilt_Pack",
)

MAX_PAGES = 40
MAX_CHARS = 120_000
MAX_FILE_MB = 25
FILE_TIMEOUT_SEC = 90
GLOSSARY_MAX_CHARS = 800_000
EXCLUDED_NAME_PATTERN = re.compile(
    r"\b(qbank|q bank|exam|questions?|answers?|mock|revision kit|practice & revision)\b",
    flags=re.I,
)
# Official/open + Merixa only: block pirate mirrors and commercial study systems
# unless explicitly licensed into the allowlist later.
BLOCKED_SOURCE_PATTERN = re.compile(
    r"\b(anna'?s?\s*archive|z-?lib|zlib\.org|libgen|sci-hub|bpp publishing|"
    r"kaplan publishing|kaplan|wiley|cia learning system|garp\s*frm|"
    r"pirate|torrent|annas-archive)\b",
    flags=re.I,
)


def clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = text.replace("\u00ad", "")
    text = (
        text.replace("\ufb01", "fi")
        .replace("\ufb02", "fl")
        .replace("\ufb00", "ff")
        .replace("\ufb03", "ffi")
        .replace("\ufb04", "ffl")
        .replace("\u201c", '"')
        .replace("\u201d", '"')
        .replace("\u2018", "'")
        .replace("\u2019", "'")
    )
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def cache_name(path: Path) -> str:
    digest = hashlib.sha1(str(path).encode("utf-8")).hexdigest()[:12]
    stem = re.sub(r"[^a-zA-Z0-9]+", "-", path.stem).strip("-")[:48] or "doc"
    return f"{stem}-{digest}.txt"


def is_priority(rel: str) -> bool:
    return any(rel.startswith(prefix) for prefix in PRIORITY_PREFIXES)


def walk(root: Path):
    if not root.exists():
        return
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        if any(part in SKIP_DIRS for part in path.parts):
            continue
        if path.name.startswith("~$"):
            continue
        if path.suffix.lower() not in {".pdf", ".docx", ".txt", ".md", ".html", ".htm"}:
            continue
        if "companies_house" in path.name.lower():
            continue
        if EXCLUDED_NAME_PATTERN.search(path.stem):
            continue
        rel_probe = str(path).replace("\\", "/")
        if BLOCKED_SOURCE_PATTERN.search(rel_probe):
            continue
        yield path


def extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path), strict=False)
    parts: list[str] = []
    for page in reader.pages[:MAX_PAGES]:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
        if sum(len(part) for part in parts) >= MAX_CHARS:
            break
    return clean("\n\n".join(parts))[:MAX_CHARS]


def extract_docx(path: Path) -> str:
    document = Document(str(path))
    parts: list[str] = [
        paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
    ]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if not cells:
                continue
            # Deduplicate merged-cell repeats while keeping order.
            unique: list[str] = []
            for cell in cells:
                if not unique or unique[-1] != cell:
                    unique.append(cell)
            parts.append("\n".join(unique))
    limit = GLOSSARY_MAX_CHARS if "glossary" in path.name.lower() else MAX_CHARS
    return clean("\n\n".join(parts))[:limit]


def extract_plain(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    if path.suffix.lower() in {".html", ".htm"}:
        raw = re.sub(r"<script[\s\S]*?</script>", " ", raw, flags=re.I)
        raw = re.sub(r"<style[\s\S]*?</style>", " ", raw, flags=re.I)
        raw = re.sub(r"</(p|div|h[1-6]|li|br|section|article)>", "\n", raw, flags=re.I)
        raw = re.sub(r"<[^>]+>", " ", raw)
    return clean(raw)[:MAX_CHARS]


def extract_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    return extract_plain(path)


def main() -> int:
    if CACHE.exists():
        for stale in CACHE.glob("*.txt"):
            stale.unlink()
    CACHE.mkdir(parents=True, exist_ok=True)

    report = {
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "extracted": 0,
        "skipped": [],
        "errors": [],
        "files": [],
    }

    for root in SOURCE_ROOTS:
        for path in walk(root):
            rel = str(path.relative_to(DOWNLOADS)).replace("\\", "/")
            if not is_priority(rel):
                continue
            if BLOCKED_SOURCE_PATTERN.search(rel):
                report["skipped"].append({"path": rel, "reason": "blocked-source"})
                print(f"skip-blocked {rel}", flush=True)
                continue

            size_mb = path.stat().st_size / (1024 * 1024)
            if size_mb > MAX_FILE_MB:
                report["skipped"].append(
                    {"path": rel, "reason": "too-large", "mb": round(size_mb, 1)}
                )
                print(f"skip-large {rel}", flush=True)
                continue

            out = CACHE / cache_name(path)
            try:
                with ThreadPoolExecutor(max_workers=1) as pool:
                    text = pool.submit(extract_file, path).result(
                        timeout=FILE_TIMEOUT_SEC
                    )

                if len(text) < 120:
                    report["skipped"].append(
                        {"path": rel, "reason": "empty-or-unreadable"}
                    )
                    print(f"skip-empty {rel}", flush=True)
                    continue

                meta = {
                    "source": rel,
                    "titleHint": path.stem,
                    "chars": len(text),
                }
                payload = f"@@META::{json.dumps(meta, ensure_ascii=True)}\n\n{text}\n"
                out.write_text(payload, encoding="utf-8")
                report["extracted"] += 1
                report["files"].append(rel)
                print(f"ok {rel} ({len(text)} chars)", flush=True)
            except FuturesTimeout:
                report["errors"].append({"path": rel, "message": "timeout"})
                print(f"timeout {rel}", flush=True)
            except Exception as exc:  # noqa: BLE001
                report["errors"].append({"path": rel, "message": str(exc)})
                print(f"err {rel}: {exc}", flush=True)

    REPORT.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(
        f"Extract complete: extracted={report['extracted']} "
        f"errors={len(report['errors'])} skipped={len(report['skipped'])}",
        flush=True,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
